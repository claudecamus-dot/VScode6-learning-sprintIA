"""Reaudit du hub VScode5 (commit bc1e74d, dimension robustesse) : une fois
qu'un fichier `_repli.jpg` existe (laisse par un echec reseau anterieur), la
branche `elif os.path.exists(chemin_repli)` le preferait TOUJOURS au vrai
fichier, meme quand la vraie photo est ensuite disponible en cache. Corrige :
la photo reelle en cache prime toujours sur le repli ; le repli ne sert que
si aucune photo reelle n'est disponible.
"""
import importlib.util
import os

import pytest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _charger(chemin):
    spec = importlib.util.spec_from_file_location("g_repli_prime_hub", chemin)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("chemin", [
    os.path.join(RACINE, "docs", "run-ia", "atelier-agentic-run", "generer_deck.py"),
])
def test_photo_reelle_en_cache_prime_sur_repli_existant(chemin, tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches

    g = _charger(chemin)
    monkeypatch.setattr(g, "IMG_DIR", str(tmp_path))

    # place_image_in_frame est importe DANS photo() (meme pattern que
    # stock_images, sys.path.insert differe) : on patche le module reel
    # apres l'avoir importe via le meme chemin.
    import sys
    sys.path.insert(0, os.path.join(RACINE, ".claude", "skills",
                                     "pptx-framed-image", "scripts"))
    import framed_image

    appels = []
    monkeypatch.setattr(
        framed_image, "place_image_in_frame",
        lambda slide, chemin_img, *a, **k: appels.append(chemin_img))

    requete = "requete-de-test-prime"
    seed = 999
    x, y, w, h = 1, 1, 2, 1
    # photo() calcule px_w=960 (fixe) et px_h a partir de l'aspect w/h ;
    # le nom de fichier cache doit correspondre exactement pour que la
    # branche "if not os.path.exists(chemin)" le trouve.
    px_w = 960
    aspect = w / h
    px_h = int(round(px_w / aspect))
    slug = "requete-de-test-prime"
    repli = "mountains"

    chemin_reel = os.path.join(
        str(tmp_path), "%s_%d_%dx%d.jpg" % (slug, seed, px_w, px_h))
    chemin_repli = os.path.join(
        str(tmp_path), "%s_%d_%dx%d_repli.jpg" % (repli, seed, px_w, px_h))

    # Les deux fichiers existent : la vraie photo en cache ET un repli
    # laisse par un echec reseau anterieur.
    with open(chemin_reel, "wb") as f:
        f.write(b"vraie-photo")
    with open(chemin_repli, "wb") as f:
        f.write(b"repli-procedural")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    g.photo(slide, x, y, w, h, requete, repli, seed=seed)

    assert appels, "place_image_in_frame n'a pas ete appele"
    assert appels[0] == chemin_reel, (
        "le repli a ete pose alors que la vraie photo etait en cache : %r"
        % appels[0])
