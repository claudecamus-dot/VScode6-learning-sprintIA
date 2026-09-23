"""Constat d'audit du hub VScode5 (dimension robustesse, #4) : le repli
image avalait toute exception (`except Exception`) et ne journalisait qu'un
`print`, indistinguable entre panne reseau et bug de crop.

Le repli reste volontairement large (offline-first : le kw-word Openverse ne
doit jamais faire echouer un build), mais l'echec doit desormais etre VISIBLE
— logue en WARNING avec le type d'exception et la trace (`exc_info=True`).
"""
import importlib.util
import logging
import os

import pytest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _charger(chemin):
    spec = importlib.util.spec_from_file_location("g_repli_hub", chemin)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize("chemin", [
    os.path.join(RACINE, "docs", "run-ia", "atelier-agentic-run", "generer_deck.py"),
])
def test_le_repli_openverse_journalise_un_warning_avec_le_type_d_exception(
        chemin, tmp_path, monkeypatch, caplog):
    from pptx import Presentation
    from pptx.util import Inches

    g = _charger(chemin)
    monkeypatch.setattr(g, "IMG_DIR", str(tmp_path))

    def _fetch_qui_echoue(dest, requete, seed=0):
        raise RuntimeError("Openverse indisponible (test)")

    # stock_images est importe DANS photo() (sys.path.insert differe) :
    # on patche apres un premier chargement pour recuperer le vrai module.
    import sys
    sys.path.insert(0, os.path.join(RACINE, ".claude", "skills",
                                     "pptx-framed-image", "scripts"))
    import stock_images
    monkeypatch.setattr(stock_images, "fetch_to", _fetch_qui_echoue)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    with caplog.at_level(logging.WARNING):
        g.photo(slide, 1, 1, 2, 1, "requete-de-test-hub", "mountains", seed=999)

    warnings = [r for r in caplog.records if r.levelno == logging.WARNING]
    assert warnings, "aucun warning journalise sur l'echec Openverse"
    assert any("RuntimeError" in r.message for r in warnings), \
        "le type d'exception n'apparait pas dans le log"
    assert any(r.exc_info for r in warnings), \
        "aucune trace (exc_info) journalisee sur l'echec"
